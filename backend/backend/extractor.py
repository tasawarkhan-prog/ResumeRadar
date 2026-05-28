import io
from pathlib import Path


def extract_text_from_pdf(file_bytes: bytes) -> str:
    try:
        import pdfplumber
        with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
            pages = [page.extract_text() or "" for page in pdf.pages]
        text = "\n".join(pages).strip()
        if text:
            return text
    except Exception:
        pass

    try:
        import PyPDF2
        reader = PyPDF2.PdfReader(io.BytesIO(file_bytes))
        text = "\n".join(p.extract_text() or "" for p in reader.pages).strip()
        if text:
            return text
    except Exception:
        pass

    raise ValueError("Could not extract text from PDF. Ensure the PDF contains selectable text (not a scanned image).")


def extract_text_from_docx(file_bytes: bytes) -> str:
    try:
        import docx
        doc = docx.Document(io.BytesIO(file_bytes))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        parts.append(cell.text.strip())
        return "\n".join(parts)
    except Exception as e:
        raise ValueError(f"Could not extract text from DOCX: {e}")


def extract_text(file_bytes: bytes, filename: str) -> str:
    ext = Path(filename).suffix.lower()
    if ext == ".pdf":
        return extract_text_from_pdf(file_bytes)
    elif ext in (".docx", ".doc"):
        return extract_text_from_docx(file_bytes)
    elif ext == ".txt":
        return file_bytes.decode("utf-8", errors="ignore").strip()
    else:
        raise ValueError(f"Unsupported file type: {ext}. Please upload a PDF, DOCX, or TXT file.")
